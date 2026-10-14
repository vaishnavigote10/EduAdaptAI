"""Multimodal ingestion: PDF pages, slides, timestamped transcripts -> located chunks -> topics -> questions."""
import re
from collections import Counter
from sklearn.cluster import KMeans
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from ..models import Chunk, Topic, Question, Document
from . import questions as qgen

TS = re.compile(r"(\d{1,2}):(\d{2})(?::(\d{2}))?[.,]?\d*\s*(?:-->\s*[\d:.,]+)?")


def kind_of(name):
    n = name.lower()
    if n.endswith(".pdf"): return "pdf"
    if n.endswith(".docx"): return "docx"
    if n.endswith((".pptx", ".ppt")): return "pptx"
    if n.endswith((".vtt", ".srt", ".mp4", ".mp3", ".wav", ".mov", ".mkv")): return "video"
    if n.endswith((".png", ".jpg", ".jpeg", ".webp")): return "image"
    return "text"


def extract_units(path, kind):
    """-> list of (location, text)."""
    units = []
    if kind == "pdf":
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(path).pages, 1):
            units.append((f"Page {i}", page.extract_text() or ""))
    elif kind == "pptx":
        from pptx import Presentation
        for i, slide in enumerate(Presentation(path).slides, 1):
            parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            units.append((f"Slide {i}", ". ".join(p for p in parts if p.strip())))
    elif kind == "docx":
        import zipfile, html as _h
        xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf8", "ignore")
        paras = [_h.unescape("".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, re.S))) for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)]
        paras = [p for p in paras if p.strip()]
        units = [(f"Section {i // 4 + 1}", " ".join(paras[i:i + 4])) for i in range(0, len(paras), 4)]
    elif kind == "video":
        if path.lower().endswith((".mp4", ".mp3", ".wav", ".mov", ".mkv")):
            try:
                import whisper  # optional: pip install openai-whisper
                res = whisper.load_model("base").transcribe(path)
                for seg in res["segments"]:
                    s = int(seg["start"]); units.append((f"{s//60}:{s%60:02d}", seg["text"]))
            except ImportError:
                raise ValueError("Audio/video needs speech-to-text. Install openai-whisper, or upload the .vtt/.srt transcript.")
        else:
            units = parse_transcript(open(path, encoding="utf-8", errors="ignore").read())
    elif kind == "image":
        try:
            import pytesseract
            from PIL import Image
            units.append(("Image", pytesseract.image_to_string(Image.open(path))))
        except ImportError:
            raise ValueError("Image text needs pytesseract + Pillow installed.")
    else:
        text = open(path, encoding="utf-8", errors="ignore").read()
        units = [(f"Section {i}", p) for i, p in enumerate([p for p in re.split(r"\n\s*\n", text) if p.strip()], 1)]
    return [(loc, re.sub(r"\s+", " ", t).strip()) for loc, t in units if len(t.strip()) > 40]


def parse_transcript(raw):
    """VTT/SRT or '[mm:ss] text' lines -> ~45s blocks keyed by start time."""
    blocks, cur_start, cur, last = [], None, [], 0
    def secs(m): return int(m.group(1)) * 60 + int(m.group(2)) if m.group(3) is None else int(m.group(1)) * 3600 + int(m.group(2)) * 60 + int(m.group(3))
    for line in raw.splitlines():
        line = line.strip()
        if not line or line.upper().startswith("WEBVTT") or line.isdigit():
            continue
        m = TS.match(line)
        if m and ("-->" in line or line.startswith("[") or len(line) < 25):
            last = secs(m)
            if cur_start is None: cur_start = last
            elif last - cur_start >= 45 and cur:
                blocks.append((cur_start, " ".join(cur))); cur, cur_start = [], last
            continue
        m2 = re.match(r"\[(\d{1,2}):(\d{2})\]\s*(.*)", line)
        if m2:
            last = int(m2.group(1)) * 60 + int(m2.group(2)); line = m2.group(3)
            if cur_start is None: cur_start = last
            elif last - cur_start >= 45 and cur:
                blocks.append((cur_start, " ".join(cur))); cur, cur_start = [], last
        cur.append(line)
    if cur:
        blocks.append((cur_start or 0, " ".join(cur)))
    return [(f"{s//60}:{s%60:02d}", t) for s, t in blocks]


def split_long(text, limit=900):
    if len(text) <= limit: return [text]
    out, cur = [], ""
    for s in re.split(r"(?<=[.!?])\s+", text):
        if len(cur) + len(s) > limit and cur: out.append(cur); cur = ""
        cur += (" " if cur else "") + s
    return out + ([cur] if cur else [])


def ingest_document(course, upload):
    doc = Document.objects.create(course=course, title=upload.name.rsplit(".", 1)[0], kind=kind_of(upload.name), file=upload)
    units = extract_units(doc.file.path, doc.kind)
    n = 0
    for loc, text in units:
        for piece in split_long(text):
            Chunk.objects.create(document=doc, text=piece, location=loc); n += 1
    if n == 0:
        doc.delete(); raise ValueError("No selectable text found. If this is a scanned PDF or an image-only file, upload a text version or paste the notes instead.")
    organise_topics(course)
    generate_questions(course)
    return doc, n


def top_terms(texts, n=2):
    c = Counter(w for t in texts for w in set(qgen.words(t)))
    return [w for w, _ in c.most_common(n)]


def _make_topics(course, chunks, X):
    k = 1 if len(chunks) < 4 else max(2, min(6, len(chunks) // 3))
    labels = [0] * len(chunks) if k == 1 else KMeans(n_clusters=k, n_init=10, random_state=3).fit_predict(X)
    start = course.topics.count(); prev = course.topics.order_by("order").last(); made = {}
    for lab in sorted(set(labels)):
        members = [chunks[i] for i in range(len(chunks)) if labels[i] == lab]
        name = " & ".join(w.title() for w in top_terms([m.text for m in members])) or f"Topic {start + len(made) + 1}"
        t = Topic.objects.create(course=course, name=name, order=start + len(made))
        if prev: t.prerequisites.add(prev)
        prev, made[lab] = t, t
    for i, c in enumerate(chunks):
        c.topic = made[labels[i]]; c.save(update_fields=["topic"])


def organise_topics(course):
    """First upload builds topics by clustering. Later uploads join the closest topic, or form new ones."""
    chunks = list(Chunk.objects.filter(document__course=course, topic__isnull=True))
    if not chunks: return
    texts = [c.text for c in chunks]
    topics = [t for t in course.topics.all() if t.chunks.exists()]
    if not topics:
        _make_topics(course, chunks, TfidfVectorizer(stop_words="english", sublinear_tf=True).fit_transform(texts)); return
    vec = TfidfVectorizer(stop_words="english", sublinear_tf=True)
    vec.fit([c.text for c in Chunk.objects.filter(document__course=course)])
    cent = vec.transform([" ".join(c.text for c in t.chunks.all()) for t in topics])
    X = vec.transform(texts); sims = cosine_similarity(X, cent); left = []
    for i, c in enumerate(chunks):
        if sims[i].max() >= 0.08:
            c.topic = topics[int(sims[i].argmax())]; c.save(update_fields=["topic"])
        else: left.append(i)
    if left:
        _make_topics(course, [chunks[i] for i in left], X[left])


def generate_questions(course):
    topics = list(course.topics.all())
    terms = {t.id: qgen.concept_terms([c.text for c in t.chunks.all()]) for t in topics}
    for t in topics:
        others = [w for o in topics if o.id != t.id for w in terms[o.id]]
        existing = list(Question.objects.filter(topic__course=course).values_list("prompt", flat=True))
        cs = [{"text": c.text, "doc": c.document.title, "location": c.location} for c in t.chunks.select_related("document")]
        for q in qgen.generate(cs, others, existing, limit=6):
            Question.objects.create(topic=t, kind="mcq", fingerprint=qgen.fingerprint(q["prompt"]), **q)
